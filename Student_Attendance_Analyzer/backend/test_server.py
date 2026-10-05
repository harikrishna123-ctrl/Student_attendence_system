"""
Live Integration Test for Student Attendance Analyzer
"""
import urllib.request
import urllib.error
import json
import uuid
from io import BytesIO
from openpyxl import load_workbook

base = "http://127.0.0.1:5000"

def test_get(endpoint):
    req = urllib.request.Request(base + endpoint)
    with urllib.request.urlopen(req) as resp:
        data = resp.read()
        print(f"[OK] {endpoint:25} -> HTTP {resp.status}, {len(data)} bytes")
        if endpoint.startswith("/api/"):
            return json.loads(data.decode("utf-8"))
        return data

def api_call(endpoint, method="GET", payload=None):
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(
        base + endpoint,
        data=body,
        headers={"Content-Type": "application/json"} if body is not None else {},
        method=method
    )
    try:
        with urllib.request.urlopen(req) as resp:
            return resp.status, resp.read()
    except urllib.error.HTTPError as err:
        return err.code, err.read()

# Test HTML Pages
test_get("/")
test_get("/students")
test_get("/attendance")
test_get("/subjects")
test_get("/analysis")
test_get("/reports")

# Test API Endpoints
status = test_get("/api/status")
print(f"  -> DB Engine: {status['database']['engine']}, Total Students: {status['total_students']}")

students = test_get("/api/students")
print(f"  -> First Student: {students['students'][0]['name']} ({students['students'][0]['student_id']})")

attendance = test_get("/api/attendance")
print(f"  -> Date: {attendance['date']} | Present: {attendance['present_count']} | Absent: {attendance['absent_count']}")

analysis = test_get("/api/analysis")
sets = analysis["sets"]
print(f"  -> Set A (Present): {sets['A']['count']}")
print(f"  -> Set B (Absent): {sets['B']['count']}")
print(f"  -> Union (A U B): {sets['union']['count']}")
print(f"  -> Intersection (A ^ B): {sets['intersection']['count']}")
print(f"  -> Diff (A - B): {sets['diff_A_minus_B']['count']}")
print(f"  -> Diff (B - A): {sets['diff_B_minus_A']['count']}")
print(f"  -> Complement A' (Absent): {sets['complement_A']['count']}")

defaulters = test_get("/api/reports/defaulters")
print(f"  -> Defaulters (<75%): {defaulters['total_defaulters']} students")

# Exercise the full workflow with temporary rows and always remove them afterward.
suffix = uuid.uuid4().hex[:10].upper()
student_id = f"IT{suffix}"
roll_number = f"IT-{suffix}"
subject_id = f"IT{suffix}"
subject_code = f"IT{suffix}"
second_subject_id = f"IT2{suffix}"
second_subject_code = f"IT2{suffix}"
student_created = False
subject_created = False
second_subject_created = False
try:
    cohort = students["students"][0]
    subject_payload = {
        "subject_id": subject_id,
        "subject_code": subject_code,
        "subject_name": "Integration Test Subject",
        "class_name": cohort["class_name"],
        "section": cohort["section"]
    }
    status_code, body = api_call("/api/subjects", "POST", subject_payload)
    assert status_code == 201, body.decode("utf-8")
    subject_created = True

    second_subject_payload = {
        "subject_id": second_subject_id,
        "subject_code": second_subject_code,
        "subject_name": "Integration Test Subject 2",
        "class_name": cohort["class_name"],
        "section": cohort["section"]
    }
    status_code, body = api_call("/api/subjects", "POST", second_subject_payload)
    assert status_code == 201, body.decode("utf-8")
    second_subject_created = True

    student_payload = {
        "student_id": student_id,
        "name": "Integration Test Student",
        "roll_number": roll_number,
        "class_name": cohort["class_name"],
        "section": cohort["section"]
    }
    status_code, body = api_call("/api/students", "POST", student_payload)
    assert status_code == 201, body.decode("utf-8")
    student_created = True

    attendance_payload = {
        "date": "2026-09-29",
        "subject_id": subject_id,
        "records": [{"student_id": student_id, "status": "Present"}]
    }
    status_code, body = api_call("/api/attendance", "POST", attendance_payload)
    assert status_code == 200, body.decode("utf-8")

    second_subject_attendance = {
        "date": attendance_payload["date"],
        "subject_id": second_subject_id,
        "records": [{"student_id": student_id, "status": "Absent"}]
    }
    status_code, body = api_call("/api/attendance", "POST", second_subject_attendance)
    assert status_code == 200, body.decode("utf-8")

    query = f"?date=2026-09-29&subject_id={subject_id}"
    status_code, body = api_call("/api/analysis" + query)
    analysis_result = json.loads(body.decode("utf-8"))
    assert status_code == 200 and student_id in {s["student_id"] for s in analysis_result["sets"]["A"]["students"]}

    attendance_payload["records"][0]["status"] = "Absent"
    status_code, body = api_call("/api/attendance", "POST", attendance_payload)
    assert status_code == 200, body.decode("utf-8")
    status_code, body = api_call("/api/attendance" + query)
    attendance_result = json.loads(body.decode("utf-8"))
    marked = [r for r in attendance_result["records"] if r["student_id"] == student_id]
    assert status_code == 200 and len(marked) == 1 and marked[0]["status"] == "Absent"

    report_query = f"?student_id={student_id}&subject_id={subject_id}&date_from=2026-09-29&date_to=2026-09-29"
    status_code, body = api_call("/api/reports/academic" + report_query)
    report = json.loads(body.decode("utf-8"))
    assert status_code == 200 and len(report["report"]) == 1
    assert report["report"][0]["total_classes"] == 1 and report["report"][0]["absent"] == 1

    status_code, body = api_call("/api/reports/export-excel" + report_query)
    workbook = load_workbook(BytesIO(body), read_only=True)
    assert status_code == 200 and workbook.active.max_row >= 5
    assert workbook.active.cell(row=5, column=2).value == student_id
    print("[OK] Student + Subject + Attendance + Analysis + Report + Excel workflow")
finally:
    if second_subject_created:
        api_call(f"/api/subjects/{second_subject_id}", "DELETE")
    if subject_created:
        api_call(f"/api/subjects/{subject_id}", "DELETE")
    if student_created:
        api_call(f"/api/students/{student_id}", "DELETE")

print("\n>>> ALL INTEGRATION TESTS PASSED 100% SUCCESSFUL! <<<")
